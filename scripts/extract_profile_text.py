#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".log",
    ".json",
    ".yaml",
    ".yml",
    ".csv",
    ".tsv",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}


class ExtractError(RuntimeError):
    pass


def compact_text(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", value)
    return re.sub(r"\n{3,}", "\n\n", normalized.replace("\r\n", "\n").strip())


def read_text_file(path: Path) -> str:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1254", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError as exc:
            last_error = exc
            continue
    raise ExtractError(f"Metin dosyasi okunamadi: {last_error}")


def extract_docx(path: Path) -> str:
    try:
        from docx import Document  # type: ignore

        doc = Document(str(path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        if paragraphs:
            return "\n".join(paragraphs)
    except Exception:
        # XML fallback below
        pass

    try:
        with zipfile.ZipFile(path) as zf:
            raw = zf.read("word/document.xml")
        root = ET.fromstring(raw)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        lines: list[str] = []
        for para in root.findall(".//w:p", ns):
            chunks = [node.text or "" for node in para.findall(".//w:t", ns)]
            merged = "".join(chunks).strip()
            if merged:
                lines.append(merged)
        return "\n".join(lines)
    except Exception as exc:  # pragma: no cover
        raise ExtractError(f"DOCX okunamadi: {exc}") from exc


def extract_pdf_text(path: Path) -> tuple[str, dict[str, Any]]:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception as exc:
        raise ExtractError(
            "PDF metin cikarma icin pypdf kurulu degil. "
            "Calistirin: pip install pypdf"
        ) from exc

    reader = PdfReader(str(path))
    lines: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            lines.append(text)
    return "\n\n".join(lines), {"page_count": len(reader.pages)}


def build_rapid_ocr():
    try:
        from rapidocr_onnxruntime import RapidOCR  # type: ignore
    except Exception as exc:
        raise ExtractError(
            "OCR icin rapidocr-onnxruntime kurulu degil. "
            "Calistirin: pip install rapidocr-onnxruntime"
        ) from exc
    return RapidOCR()


def ocr_image(path: Path) -> tuple[str, dict[str, Any]]:
    ocr = build_rapid_ocr()
    result, _elapsed = ocr(str(path))
    lines = [item[1].strip() for item in (result or []) if len(item) > 1 and str(item[1]).strip()]
    return "\n".join(lines), {"line_count": len(lines)}


def ocr_pdf(path: Path, page_limit: int = 20, render_scale: float = 2.0) -> tuple[str, dict[str, Any]]:
    ocr = build_rapid_ocr()

    try:
        import numpy as np  # type: ignore
    except Exception as exc:
        raise ExtractError(
            "OCR icin numpy kurulu degil. Calistirin: pip install numpy"
        ) from exc

    try:
        import pypdfium2 as pdfium  # type: ignore
    except Exception as exc:
        raise ExtractError(
            "PDF OCR icin pypdfium2 kurulu degil. Calistirin: pip install pypdfium2"
        ) from exc

    try:
        doc = pdfium.PdfDocument(str(path))
    except Exception as exc:
        raise ExtractError(f"PDF OCR icin dosya acilamadi: {exc}") from exc

    total_pages = len(doc)
    pages_to_read = min(total_pages, max(1, page_limit))
    parts: list[str] = []

    for idx in range(pages_to_read):
        page = doc[idx]
        bitmap = page.render(scale=render_scale)
        pil_img = bitmap.to_pil()
        arr = np.array(pil_img)
        result, _elapsed = ocr(arr)
        text_lines = [item[1].strip() for item in (result or []) if len(item) > 1 and str(item[1]).strip()]
        text = "\n".join(text_lines).strip()
        if text:
            parts.append(f"[page {idx + 1}]\n{text}")
        try:
            page.close()
        except Exception:
            pass

    return "\n\n".join(parts), {
        "page_count": total_pages,
        "ocr_pages": pages_to_read,
    }


def extract(path: Path, ocr_mode: str, max_chars: int) -> dict[str, Any]:
    if not path.exists() or not path.is_file():
        raise ExtractError("Dosya bulunamadi.")

    ext = path.suffix.lower()
    warnings: list[str] = []
    parser = "text"
    ocr_used = False
    metadata: dict[str, Any] = {}

    if ext in TEXT_EXTENSIONS:
        text = read_text_file(path)
        parser = "text"
    elif ext == ".docx":
        text = extract_docx(path)
        parser = "docx"
    elif ext == ".pdf":
        text, pdf_meta = extract_pdf_text(path)
        metadata.update(pdf_meta)
        parser = "pdf_text"

        should_ocr = ocr_mode == "always" or (ocr_mode == "auto" and len(text.strip()) < 80)
        if should_ocr:
            ocr_text, ocr_meta = ocr_pdf(path)
            metadata.update(ocr_meta)
            if ocr_text.strip():
                if len(text.strip()) < 80 or len(ocr_text) > len(text) * 1.1:
                    text = ocr_text
                    parser = "ocr"
                    ocr_used = True
                else:
                    warnings.append("OCR sonucu mevcut PDF metninden daha iyi gorunmedi, PDF metni kullanildi.")
            elif ocr_mode == "always":
                raise ExtractError("OCR zorunlu secildi ama OCR metni cikmadi.")
    elif ext in IMAGE_EXTENSIONS:
        if ocr_mode == "never":
            raise ExtractError("Gorsel dosyalarda OCR kapali iken metin cikarimi yapilamaz.")
        text, ocr_meta = ocr_image(path)
        metadata.update(ocr_meta)
        parser = "ocr"
        ocr_used = True
    else:
        raise ExtractError(
            "Desteklenmeyen dosya uzantisi. Desteklenenler: "
            "pdf, docx, txt, md, json, yaml, csv, tsv, png, jpg, jpeg, webp, bmp."
        )

    text = compact_text(text)
    if not text:
        raise ExtractError("Dosyadan anlamli metin cikmadi.")

    if len(text) > max_chars:
        text = text[:max_chars]
        warnings.append(f"Metin {max_chars} karakter ile sinirlandi.")

    return {
        "ok": True,
        "path": str(path),
        "parser": parser,
        "ocr_used": ocr_used,
        "text": text,
        "text_chars": len(text),
        "warnings": warnings,
        "metadata": metadata,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Extract text from profile files with optional OCR")
    parser.add_argument("--path", required=True, help="Input file path")
    parser.add_argument("--ocr", choices=("auto", "always", "never"), default="auto")
    parser.add_argument("--max-chars", type=int, default=90000)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    try:
        payload = extract(Path(args.path).resolve(), args.ocr, args.max_chars)
        print(json.dumps(payload, ensure_ascii=True))
        return 0
    except ExtractError as exc:
        print(
            json.dumps(
                {
                    "ok": False,
                    "error": str(exc),
                },
                ensure_ascii=True,
            )
        )
        return 2
    except Exception as exc:  # pragma: no cover
        print(
            json.dumps(
                {
                    "ok": False,
                    "error": f"Beklenmeyen hata: {exc}",
                },
                ensure_ascii=True,
            )
        )
        return 3


if __name__ == "__main__":
    sys.exit(main())
